"""C-485 — stream the tracked DSLD XLSX workbooks into the local build DB.

The source is intentionally read-only under docs/ssot/daten/.  CSV is only
the PostgreSQL COPY wire format; it is never written as a second 2m-row file.
"""
from __future__ import annotations

import argparse
import csv
import io
import json
import re
import subprocess
import sys
import time
import uuid
from collections import defaultdict
from pathlib import Path

import openpyxl


CONTAINER = 'supabase_db_LumeOS-Claude-V1'
ROLE_COLUMNS = {
    'Manufacturer': 'manufacturer',
    'Distributor': 'distributor',
    'Packager': 'packager',
    'Reseller': 'reseller',
    'Other': 'other',
}


def text(value: object) -> str:
    return '' if value is None else str(value).strip()


def normalized(value: object) -> str:
    return re.sub(r'\s+', ' ', text(value)).casefold()


def yes(value: object) -> bool:
    return normalized(value) == 'yes'


def numeric_and_unit(value: object) -> tuple[str, str]:
    match = re.match(r'^\s*(\d+(?:\.\d+)?)\s+(.+?)\s*$', text(value))
    return (match.group(1), match.group(2)) if match else ('', '')


def barcode(value: object) -> str:
    digits = ''.join(re.findall(r'\d', text(value)))
    return digits if 8 <= len(digits) <= 14 else ''


def amount(value: object) -> tuple[str, str, str]:
    raw = text(value)
    if not raw:
        return ('', 'not_stated', '')
    match = re.fullmatch(r'([<>])\s*(\d+(?:[.,]\d+)?)', raw)
    if match:
        return (match.group(2).replace(',', '.'), 'less_than' if match.group(1) == '<' else 'greater_than', raw)
    if re.fullmatch(r'\d+(?:[.,]\d+)?', raw):
        return (raw.replace(',', '.'), 'exact', raw)
    return ('', 'not_stated', raw)


def workbooks(source: Path):
    return sorted(source.glob('xlsx-batch*.xlsx'), key=lambda item: int(re.search(r'(\d+)', item.stem).group(1)))


def rows(book: openpyxl.Workbook, sheet_name: str):
    sheet = book[sheet_name]
    iterator = sheet.iter_rows(values_only=True)
    headers = [text(value) for value in next(iterator)]
    for values in iterator:
        yield dict(zip(headers, values))


def psql(database: str, sql: str) -> str:
    result = subprocess.run(
        ['docker', 'exec', CONTAINER, 'psql', '-X', '-q', '-v', 'ON_ERROR_STOP=1', '-U', 'postgres', '-d', database, '-t', '-A', '-c', sql],
        check=True, capture_output=True, text=True,
    )
    return result.stdout.strip()


def copy(database: str, before: str, columns: list[str], data, after: str) -> None:
    process = subprocess.Popen(
        ['docker', 'exec', '-i', CONTAINER, 'psql', '-X', '-q', '-v', 'ON_ERROR_STOP=1', '-U', 'postgres', '-d', database, '-f', '-'],
        stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
        text=True, encoding='utf-8', errors='strict',
    )
    assert process.stdin is not None
    process.stdin.write(before)
    process.stdin.write(f"COPY tmp_c485 ({','.join(columns)}) FROM STDIN WITH (FORMAT csv);\n")
    writer = csv.writer(process.stdin, lineterminator='\n')
    for row in data:
        writer.writerow(row)
    process.stdin.write('\\.\n')
    process.stdin.write(after)
    process.stdin.close()
    stdout = process.stdout.read() if process.stdout else ''
    stderr = process.stderr.read() if process.stderr else ''
    status = process.wait()
    if stdout:
        print(stdout, end='')
    if status:
        raise RuntimeError(stderr.strip() or f'COPY schlug fehl ({status})')
    if stderr:
        print(stderr, file=sys.stderr, end='')


def product_rows(files):
    for file in files:
        book = openpyxl.load_workbook(file, read_only=True, data_only=True)
        for row in rows(book, 'Product Overview'):
            dsld_id = text(row['DSLD ID'])
            if not dsld_id:
                continue
            net_size, net_unit = numeric_and_unit(row['Net Contents'])
            serving_size, serving_unit = numeric_and_unit(row['Serving Size'])
            status = text(row['Market Status'])
            yield (
                dsld_id, text(row['Product Name']), text(row['Brand Name']), barcode(row['Bar Code']),
                net_size, net_unit, serving_size, serving_unit,
                text(row['Product Type [LanguaL]']), text(row['Supplement Form [LanguaL]']),
                text(row['Date Entered into DSLD']), status,
                'true' if normalized(status) == 'on market' else 'false', text(row['Suggested Use']),
            )


def company_rows(files):
    for file in files:
        book = openpyxl.load_workbook(file, read_only=True, data_only=True)
        for row in rows(book, 'Company Information'):
            name = text(row['Company Name'])
            dsld_id = text(row['DSLD ID'])
            if not name or not dsld_id:
                continue
            for field, role in ROLE_COLUMNS.items():
                if yes(row[field]):
                    yield (dsld_id, name, text(row['Country']), role)


def catalog(database: str) -> dict[str, set[str]]:
    payload = psql(database, """
      SELECT coalesce(jsonb_agg(row_to_json(x)), '[]'::jsonb)::text
      FROM (
        SELECT s.id::text AS supplement_id, s.name_de, s.name_en, s.name_th, a.alias
        FROM supplements.supplements s
        LEFT JOIN supplements.supplement_aliases a ON a.supplement_id = s.id
      ) x;
    """)
    result: dict[str, set[str]] = defaultdict(set)
    for row in json.loads(payload):
        for name in (row['name_de'], row['name_en'], row['name_th'], row['alias']):
            if name:
                result[normalized(name)].add(row['supplement_id'])
    return result


def fact_rows(files, aliases: dict[str, set[str]], counts: dict[str, int], metrics: dict[str, int]):
    for file in files:
        book = openpyxl.load_workbook(file, read_only=True, data_only=True)
        active_product = ''
        active_blend = ''
        for row in rows(book, 'Dietary Supplement Facts'):
            product = text(row['DSLD ID'])
            ingredient = text(row['Ingredient'])
            if not product or not ingredient:
                continue
            if product != active_product:
                active_product = product
                active_blend = ''
            counts[product] += 1
            order = counts[product]
            category = text(row['DSLD Ingredient Categories'])
            raw_amount = text(row['Amount Per Serving'])
            amount_value, amount_qualifier, amount_raw = amount(raw_amount)
            if normalized(category) == 'blend':
                active_blend = str(order)
                blend_order = ''
            elif not raw_amount and active_blend:
                blend_order = active_blend
            else:
                active_blend = ''
                blend_order = ''
            found = aliases.get(normalized(ingredient), set())
            supplement_id = next(iter(found)) if len(found) == 1 else ''
            metrics['fact_rows'] += 1
            metrics['matched_rows'] += int(bool(supplement_id))
            metrics['candidate_rows'] += int(not supplement_id)
            yield (
                str(uuid.uuid4()), product, supplement_id, ingredient, category, amount_value, amount_qualifier,
                amount_raw, text(row['Amount Per Serving Unit']), blend_order, str(order), 'true',
            )


def other_ingredient_rows(files, counts: dict[str, int], metrics: dict[str, int]):
    for file in files:
        book = openpyxl.load_workbook(file, read_only=True, data_only=True)
        for row in rows(book, 'Other Ingredients'):
            product = text(row['DSLD ID'])
            raw = text(row['Other Ingredients'])
            if not product or not raw:
                continue
            # DSLD puts this list in one cell; csv keeps quoted commas intact.
            for ingredient in next(csv.reader(io.StringIO(raw), skipinitialspace=True)):
                ingredient = ingredient.strip()
                if not ingredient:
                    continue
                counts[product] += 1
                metrics['other_ingredient_rows'] += 1
                yield (
                    str(uuid.uuid4()), product, '', ingredient, 'other ingredient', '', 'not_stated', '', '', '',
                    str(counts[product]), 'false',
                )


def run(database: str, source: Path) -> None:
    files = workbooks(source)
    if len(files) != 11:
        raise RuntimeError(f'C-485: erwartete 11 XLSX-Dateien, gefunden: {len(files)}')
    started = time.perf_counter()
    psql(database, """
      BEGIN;
      DELETE FROM supplements.supplement_field_sources WHERE source = 'dsld';
      DELETE FROM supplements.product_content_candidates WHERE source = 'dsld';
      DELETE FROM supplements.product_contents WHERE source = 'dsld';
      DELETE FROM supplements.product_suppliers WHERE source = 'dsld';
      DELETE FROM supplements.supplier_products WHERE source = 'dsld';
      DELETE FROM supplements.suppliers s
      WHERE s.source = 'dsld'
        AND NOT EXISTS (SELECT 1 FROM supplements.product_suppliers ps WHERE ps.supplier_id = s.id)
        AND NOT EXISTS (SELECT 1 FROM supplements.supplier_products sp WHERE sp.supplier_id = s.id);
      COMMIT;
    """)

    stage_started = time.perf_counter()
    copy(database, """
      BEGIN;
      CREATE TEMP TABLE tmp_c485 (dsld_id bigint, name_en text, marke text, gtin text, net_size numeric, net_unit text, serving_size numeric, serving_unit text, product_type text, produktform text, date_entered date, market_status text, is_active boolean, suggested_use text) ON COMMIT DROP;
    """, ['dsld_id','name_en','marke','gtin','net_size','net_unit','serving_size','serving_unit','product_type','produktform','date_entered','market_status','is_active','suggested_use'], product_rows(files), """
      INSERT INTO supplements.supplier_products (name_en, marke, dsld_id, gtin, packungsgroesse, packungseinheit, portionsgroesse, portionseinheit, product_type, produktform, date_entered, market_status, is_active, im_katalog, suggested_use, source)
      SELECT name_en, nullif(marke,''), dsld_id, nullif(gtin,''), net_size, nullif(net_unit,''), serving_size, nullif(serving_unit,''), nullif(product_type,''), nullif(produktform,''), date_entered, nullif(market_status,''), is_active, true, nullif(suggested_use,''), 'dsld'
      FROM tmp_c485;
      COMMIT;
    """)
    products_seconds = time.perf_counter() - stage_started

    stage_started = time.perf_counter()
    copy(database, """
      BEGIN;
      CREATE TEMP TABLE tmp_c485 (dsld_id bigint, company_name text, country text, rolle text) ON COMMIT DROP;
    """, ['dsld_id','company_name','country','rolle'], company_rows(files), """
      INSERT INTO supplements.suppliers (name, land, source)
      SELECT min(company_name), nullif(min(country),''), 'dsld'
      FROM tmp_c485
      GROUP BY lower(regexp_replace(btrim(company_name), '\\s+', ' ', 'g'))
      ON CONFLICT (name_normalized) DO NOTHING;
      INSERT INTO supplements.product_suppliers (product_id, supplier_id, rolle, source)
      SELECT p.id, s.id, t.rolle, 'dsld'
      FROM tmp_c485 t
      JOIN supplements.supplier_products p ON p.dsld_id = t.dsld_id
      JOIN supplements.suppliers s ON s.name_normalized = lower(regexp_replace(btrim(t.company_name), '\\s+', ' ', 'g'))
      ON CONFLICT DO NOTHING;
      COMMIT;
    """)
    companies_seconds = time.perf_counter() - stage_started

    stage_started = time.perf_counter()
    aliases = catalog(database)
    counts: dict[str, int] = defaultdict(int)
    metrics = {'fact_rows': 0, 'matched_rows': 0, 'candidate_rows': 0, 'other_ingredient_rows': 0}
    def all_contents():
        yield from fact_rows(files, aliases, counts, metrics)
        yield from other_ingredient_rows(files, counts, metrics)
    copy(database, """
      BEGIN;
      CREATE TEMP TABLE tmp_c485 (id uuid, product_dsld_id bigint, supplement_id uuid, ingredient_name text, ingredient_category text, amount_per_serving numeric, amount_qualifier text, amount_raw text, unit text, blend_reihenfolge integer, reihenfolge integer, ist_wirkstoff boolean) ON COMMIT DROP;
    """, ['id','product_dsld_id','supplement_id','ingredient_name','ingredient_category','amount_per_serving','amount_qualifier','amount_raw','unit','blend_reihenfolge','reihenfolge','ist_wirkstoff'], all_contents(), """
      INSERT INTO supplements.product_contents (id, product_id, supplement_id, ingredient_name, ingredient_category, amount_per_serving, amount_qualifier, amount_raw, unit, ist_wirkstoff, source, reihenfolge)
      SELECT t.id, p.id, t.supplement_id, t.ingredient_name, nullif(t.ingredient_category,''), t.amount_per_serving, t.amount_qualifier, nullif(t.amount_raw,''), nullif(t.unit,''), t.ist_wirkstoff, 'dsld', t.reihenfolge
      FROM tmp_c485 t
      JOIN supplements.supplier_products p ON p.dsld_id = t.product_dsld_id;
      UPDATE supplements.product_contents child
      SET blend_id = parent.id
      FROM tmp_c485 t
      JOIN supplements.supplier_products p ON p.dsld_id = t.product_dsld_id
      JOIN supplements.product_contents parent ON parent.product_id = p.id AND parent.reihenfolge = t.blend_reihenfolge
      WHERE child.id = t.id AND t.blend_reihenfolge IS NOT NULL;
      INSERT INTO supplements.product_content_candidates (product_id, product_content_id, ingredient_name, amount_per_serving, amount_qualifier, amount_raw, unit, source)
      SELECT p.id, t.id, t.ingredient_name, t.amount_per_serving, t.amount_qualifier, nullif(t.amount_raw,''), nullif(t.unit,''), 'dsld'
      FROM tmp_c485 t
      JOIN supplements.supplier_products p ON p.dsld_id = t.product_dsld_id
      WHERE t.ist_wirkstoff AND t.supplement_id IS NULL;
      COMMIT;
    """)
    contents_seconds = time.perf_counter() - stage_started

    stage_started = time.perf_counter()
    psql(database, """
      BEGIN;
      INSERT INTO supplements.supplement_field_sources (supplier_product_id, status, field_name, source, source_id, as_of, evidence_class, source_note_en)
      SELECT p.id, 'bekannt', f.field_name, 'dsld', 'dsld:' || p.dsld_id::text, p.date_entered, 'dsld', 'DSLD Product Overview'
      FROM supplements.supplier_products p
      CROSS JOIN LATERAL (VALUES
        ('name_en', p.name_en), ('marke', p.marke), ('dsld_id', p.dsld_id::text), ('gtin', p.gtin),
        ('packungsgroesse', p.packungsgroesse::text), ('packungseinheit', p.packungseinheit),
        ('portionsgroesse', p.portionsgroesse::text), ('portionseinheit', p.portionseinheit),
        ('product_type', p.product_type), ('produktform', p.produktform), ('market_status', p.market_status),
        ('date_entered', p.date_entered::text), ('suggested_use', p.suggested_use)
      ) AS f(field_name, value)
      WHERE p.source = 'dsld' AND nullif(btrim(f.value),'') IS NOT NULL;
      COMMIT;
    """)
    sources_seconds = time.perf_counter() - stage_started

    summary = json.loads(psql(database, """
      SELECT json_build_object(
        'products', (SELECT count(*) FROM supplements.supplier_products WHERE source = 'dsld'),
        'companies', (SELECT count(*) FROM supplements.suppliers WHERE source = 'dsld'),
        'product_suppliers', (SELECT count(*) FROM supplements.product_suppliers WHERE source = 'dsld'),
        'contents', (SELECT count(*) FROM supplements.product_contents WHERE source = 'dsld'),
        'candidates', (SELECT count(*) FROM supplements.product_content_candidates WHERE source = 'dsld'),
        'field_sources', (SELECT count(*) FROM supplements.supplement_field_sources WHERE source = 'dsld')
      )::text;
    """))
    elapsed = time.perf_counter() - started
    print('C-485 DSLD ' + json.dumps({
        **summary,
        **metrics,
        'seconds': round(elapsed, 3),
        'products_seconds': round(products_seconds, 3),
        'companies_seconds': round(companies_seconds, 3),
        'contents_seconds': round(contents_seconds, 3),
        'sources_seconds': round(sources_seconds, 3),
    }, ensure_ascii=False))


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--database', required=True)
    parser.add_argument('--source', required=True)
    args = parser.parse_args()
    run(args.database, Path(args.source))
