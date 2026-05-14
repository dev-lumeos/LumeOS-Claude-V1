#!/usr/bin/env python
"""Deterministic local-only BLS food sample staging helper.

This helper reads the approved local BLS workbook and emits a small reviewable
SQL staging payload for the local Supabase/Test DB. It does not connect to any
database and it does not perform a broad import.
"""

from __future__ import annotations

import argparse
import re
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook


DEFAULT_WORKBOOK_PATH = Path("docs/specs/Nutrition/00_raw/bls/original/BLS_4_0_Daten_2025_DE.xlsx")
DEFAULT_NUTRIENT_DEFS_PATH = Path("docs/project/p1-005/P1-005-nutrient-defs-seed-candidate.md")
DEFAULT_SQL_OUTPUT_PATH = Path("docs/project/p1-005/P1-005-local-food-sample-staging.sql")
DEFAULT_REPORT_OUTPUT_PATH = Path("docs/project/p1-005/P1-005-local-food-sample-staging-report.md")

MISSING_MARKERS = {"", "-", "TR", "<LOQ", "<LOD"}
SOURCE_LABEL = "bls_4_0_local_sample"


@dataclass(frozen=True)
class FoodRow:
    bls_code: str
    name_de: str
    name_en: str
    source_row: int


@dataclass(frozen=True)
class FoodNutrientRow:
    bls_code: str
    nutrient_code: str
    value: Decimal
    source_row: int
    source_column: int


@dataclass(frozen=True)
class FoodSampleExtraction:
    source_workbook: Path
    nutrient_defs_path: Path
    sheet_name: str
    foods: list[FoodRow]
    nutrients: list[FoodNutrientRow]
    supported_nutrient_codes: set[str]
    skipped_unsupported_codes: set[str]
    skipped_missing_values: int


def parse_bls_value(value: object) -> float | None:
    """Parse a BLS nutrient value according to SPEC_08_IMPORT_PIPELINE."""

    if value is None:
        return None
    if isinstance(value, int | float):
        return float(value)
    text = str(value).strip()
    if text.upper() in MISSING_MARKERS:
        return None
    if text.startswith("<"):
        return None
    text = text.replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None


def parse_bls_decimal(value: object) -> Decimal | None:
    parsed = parse_bls_value(value)
    if parsed is None:
        return None
    try:
        return Decimal(str(parsed)).quantize(Decimal("0.00001"))
    except InvalidOperation:
        return None


def extract_nutrient_code(header: object) -> str | None:
    if header is None:
        return None
    text = str(header).strip()
    if not text:
        return None
    code = text.split()[0].strip()
    if re.fullmatch(r"[A-Z0-9]+", code):
        return code
    return None


def parse_supported_nutrient_codes(markdown: str) -> set[str]:
    codes: set[str] = set()
    for line in markdown.splitlines():
        stripped = line.strip()
        if not stripped.startswith("|") or "---" in stripped:
            continue
        cells = [cell.strip() for cell in stripped.strip("|").split("|")]
        if not cells:
            continue
        code = cells[0]
        if code.lower() == "code":
            continue
        if re.fullmatch(r"[A-Z0-9]+", code):
            codes.add(code)
    return codes


def sql_quote(value: str) -> str:
    return "'" + value.replace("'", "''") + "'"


def decimal_sql(value: Decimal) -> str:
    return format(value, "f")


def extract_food_sample(
    *,
    workbook_path: Path = DEFAULT_WORKBOOK_PATH,
    nutrient_defs_path: Path = DEFAULT_NUTRIENT_DEFS_PATH,
    limit: int = 10,
) -> FoodSampleExtraction:
    if limit <= 0:
        raise ValueError("limit must be positive")
    if not workbook_path.exists():
        raise FileNotFoundError(f"BLS workbook is missing: {workbook_path}")
    if not nutrient_defs_path.exists():
        raise FileNotFoundError(f"nutrient_defs candidate is missing: {nutrient_defs_path}")

    supported_codes = parse_supported_nutrient_codes(nutrient_defs_path.read_text(encoding="utf-8"))
    if not supported_codes:
        raise ValueError(f"no supported nutrient codes parsed from {nutrient_defs_path}")

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        header = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]

        nutrient_columns: list[tuple[int, str]] = []
        skipped_unsupported_codes: set[str] = set()
        for column_index in range(4, len(header) + 1, 3):
            code = extract_nutrient_code(header[column_index - 1])
            if code is None:
                continue
            if code in supported_codes:
                nutrient_columns.append((column_index, code))
            else:
                skipped_unsupported_codes.add(code)

        foods: list[FoodRow] = []
        nutrients: list[FoodNutrientRow] = []
        skipped_missing_values = 0

        for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            bls_code = str(row[0] or "").strip()
            name_de = str(row[1] or "").strip()
            name_en = str(row[2] or "").strip()
            if not bls_code or not name_de:
                continue

            foods.append(FoodRow(
                bls_code=bls_code,
                name_de=name_de,
                name_en=name_en,
                source_row=row_number,
            ))

            for column_index, nutrient_code in nutrient_columns:
                value = row[column_index - 1] if column_index - 1 < len(row) else None
                parsed = parse_bls_decimal(value)
                if parsed is None:
                    skipped_missing_values += 1
                    continue
                nutrients.append(FoodNutrientRow(
                    bls_code=bls_code,
                    nutrient_code=nutrient_code,
                    value=parsed,
                    source_row=row_number,
                    source_column=column_index,
                ))

            if len(foods) >= limit:
                break
        sheet_name = sheet.title
    finally:
        workbook.close()

    if len(foods) < limit:
        raise ValueError(f"only {len(foods)} source-backed foods found; requested {limit}")
    if not nutrients:
        raise ValueError("no source-backed nutrient values found for selected foods")

    return FoodSampleExtraction(
        source_workbook=workbook_path,
        nutrient_defs_path=nutrient_defs_path,
        sheet_name=sheet_name,
        foods=foods,
        nutrients=nutrients,
        supported_nutrient_codes=supported_codes,
        skipped_unsupported_codes=skipped_unsupported_codes,
        skipped_missing_values=skipped_missing_values,
    )


def build_values(rows: Iterable[tuple[str, ...]]) -> str:
    return ",\n    ".join("(" + ", ".join(row) + ")" for row in rows)


def build_local_staging_sql(extraction: FoodSampleExtraction) -> str:
    food_values = build_values(
        (
            sql_quote(food.bls_code),
            sql_quote(food.name_de),
            sql_quote(food.name_en),
            "''",
            sql_quote(food.name_de),
        )
        for food in extraction.foods
    )
    nutrient_values = build_values(
        (
            sql_quote(row.bls_code),
            sql_quote(row.nutrient_code),
            decimal_sql(row.value),
            sql_quote(SOURCE_LABEL),
        )
        for row in extraction.nutrients
    )

    return f"""-- STATUS: LOCAL_ONLY_FOOD_SAMPLE_STAGING
-- PURPOSE: P1-005 deterministic local food sample staging for Food Search preparation.
-- SOURCE:
--   - {extraction.source_workbook.as_posix()}
--   - sheet: {extraction.sheet_name}
--   - nutrient code allowlist: {extraction.nutrient_defs_path.as_posix()}
-- AUTHORIZATION BOUNDARY:
--   - Local Supabase/Test DB only
--   - No DEV/LIVE
--   - No Supabase Cloud
--   - No raw BLS commit
--   - No broad/full BLS import
--   - No invented food values
--   - No RDA value changes
--
-- EXPECTED LOCAL POST-APPLY VALIDATION:
--   SELECT count(*) FROM nutrition.foods;
--   SELECT count(*) FROM nutrition.food_nutrients;
--   SELECT count(*)
--   FROM nutrition.food_nutrients fn
--   LEFT JOIN nutrition.nutrient_defs nd ON nd.code = fn.nutrient_code
--   WHERE nd.code IS NULL;

begin;

do $$
begin
  if to_regclass('nutrition.foods') is null then
    raise exception 'nutrition.foods is missing; apply local food foundation schema first';
  end if;

  if to_regclass('nutrition.food_nutrients') is null then
    raise exception 'nutrition.food_nutrients is missing; apply local food foundation schema first';
  end if;

  if to_regclass('nutrition.nutrient_defs') is null then
    raise exception 'nutrition.nutrient_defs is missing; apply local nutrient foundation first';
  end if;
end
$$;

insert into nutrition.foods (bls_code, name_de, name_en, name_th, name_display)
values
    {food_values}
on conflict (bls_code) do update set
  name_de = excluded.name_de,
  name_en = excluded.name_en,
  name_th = excluded.name_th,
  name_display = excluded.name_display,
  updated_at = now();

insert into nutrition.food_nutrients (food_id, nutrient_code, value, data_source)
select
  foods.id,
  source_rows.nutrient_code,
  source_rows.value,
  source_rows.data_source
from (
  values
    {nutrient_values}
) as source_rows(bls_code, nutrient_code, value, data_source)
join nutrition.foods foods
  on foods.bls_code = source_rows.bls_code
join nutrition.nutrient_defs nutrient_defs
  on nutrient_defs.code = source_rows.nutrient_code
on conflict (food_id, nutrient_code) do update set
  value = excluded.value,
  data_source = excluded.data_source;

analyze nutrition.foods;
analyze nutrition.food_nutrients;

commit;
"""


def build_report(extraction: FoodSampleExtraction) -> str:
    food_rows = "\n".join(
        f"| {food.bls_code} | {food.name_de} | {food.name_en} | {food.source_row} |"
        for food in extraction.foods
    )
    distinct_nutrient_count = len({row.nutrient_code for row in extraction.nutrients})

    return f"""# P1-005 Local Food Sample Staging Report

Status: LOCAL_ONLY_REVIEWED_STAGING_CANDIDATE

## Purpose

Create a small source-backed local Nutrition food sample so the next local step
can build real Food Search without inventing food values or running a broad BLS
import.

## Source-backed deterministic sample

- Source workbook: `{extraction.source_workbook.as_posix()}`
- Sheet: `{extraction.sheet_name}`
- Nutrient code allowlist: `{extraction.nutrient_defs_path.as_posix()}`
- Selection rule: first {len(extraction.foods)} rows with non-empty BLS code and German food name
- Expected food rows: {len(extraction.foods)}
- Expected food_nutrients rows: {len(extraction.nutrients)}
- Distinct nutrient codes inserted: {distinct_nutrient_count}
- Skipped missing/trace marker values: {extraction.skipped_missing_values}
- Unsupported nutrient headers skipped: {len(extraction.skipped_unsupported_codes)}

## Food rows

| bls_code | name_de | name_en | source row |
|---|---|---|---:|
{food_rows}

## Boundaries

- Local Supabase/Test DB only.
- No DEV/LIVE action.
- No Supabase Cloud command.
- No raw BLS commit.
- No broad/full BLS import.
- No invented food values.
- No RDA value changes.

## Validation queries

```sql
select count(*) from nutrition.foods;
select count(*) from nutrition.food_nutrients;
select count(*)
from nutrition.food_nutrients fn
left join nutrition.nutrient_defs nd on nd.code = fn.nutrient_code
where nd.code is null;
```
"""


def write_outputs(
    *,
    extraction: FoodSampleExtraction,
    sql_output_path: Path = DEFAULT_SQL_OUTPUT_PATH,
    report_output_path: Path = DEFAULT_REPORT_OUTPUT_PATH,
) -> None:
    sql_output_path.parent.mkdir(parents=True, exist_ok=True)
    report_output_path.parent.mkdir(parents=True, exist_ok=True)
    sql_output_path.write_text(build_local_staging_sql(extraction), encoding="utf-8", newline="\n")
    report_output_path.write_text(build_report(extraction), encoding="utf-8", newline="\n")


def main() -> int:
    parser = argparse.ArgumentParser(description="Build deterministic local-only BLS food sample staging artifacts.")
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK_PATH)
    parser.add_argument("--nutrient-defs", type=Path, default=DEFAULT_NUTRIENT_DEFS_PATH)
    parser.add_argument("--sql-output", type=Path, default=DEFAULT_SQL_OUTPUT_PATH)
    parser.add_argument("--report-output", type=Path, default=DEFAULT_REPORT_OUTPUT_PATH)
    parser.add_argument("--limit", type=int, default=10)
    args = parser.parse_args()

    extraction = extract_food_sample(
        workbook_path=args.workbook,
        nutrient_defs_path=args.nutrient_defs,
        limit=args.limit,
    )
    write_outputs(
        extraction=extraction,
        sql_output_path=args.sql_output,
        report_output_path=args.report_output,
    )
    print(f"foods={len(extraction.foods)}")
    print(f"food_nutrients={len(extraction.nutrients)}")
    print(f"sql={args.sql_output}")
    print(f"report={args.report_output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
