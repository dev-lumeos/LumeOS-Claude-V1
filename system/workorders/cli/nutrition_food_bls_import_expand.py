#!/usr/bin/env python
"""Deterministic local-only BLS food import expansion helper.

The helper extracts source-backed BLS food rows and nutrient values into UTF-8
CSV files for local COPY. It intentionally keeps bulk BLS-derived CSV data in a
local temp directory instead of committing it to git.
"""

from __future__ import annotations

import argparse
import csv
import re
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook


DEFAULT_WORKBOOK_PATH = Path("docs/specs/Nutrition/00_raw/bls/original/BLS_4_0_Daten_2025_DE.xlsx")
DEFAULT_NUTRIENT_DEFS_PATH = Path("docs/project/p1-005/P1-005-nutrient-defs-seed-candidate.md")
DEFAULT_OUTPUT_DIR = Path("tmp/nutrition/p1-005-bls-local-import")
DEFAULT_REPORT_PATH = Path("docs/project/p1-005/P1-005-local-bls-import-expansion-report.md")
CONTAINER_IMPORT_DIR = "/tmp/p1-005-bls-local-import"
SOURCE_LABEL = "bls_4_0_local_import"
MISSING_MARKERS = {"", "-", "TR", "<LOQ", "<LOD"}


@dataclass(frozen=True)
class ImportResult:
    source_workbook: Path
    nutrient_defs_path: Path
    output_dir: Path
    sheet_name: str
    food_count: int
    food_nutrient_count: int
    supported_nutrient_count: int
    distinct_nutrient_count: int
    missing_value_count: int
    unsupported_nutrient_codes: set[str]
    sample_foods: list[tuple[str, str, str]]


def parse_bls_decimal(value: object) -> Decimal | None:
    if value is None:
        return None
    if isinstance(value, int | float):
        text = str(value)
    else:
        text = str(value).strip()
        if text.upper() in MISSING_MARKERS:
            return None
        if text.startswith("<"):
            return None
        text = text.replace(",", ".")
    try:
        return Decimal(text).quantize(Decimal("0.00001"))
    except (InvalidOperation, ValueError):
        return None


def extract_nutrient_code(header: object) -> str | None:
    if header is None:
        return None
    text = str(header).strip()
    if not text:
        return None
    code = text.split()[0].strip()
    return code if re.fullmatch(r"[A-Z0-9]+", code) else None


def parse_supported_nutrient_codes(markdown: str) -> set[str]:
    codes: set[str] = set()
    for line in markdown.splitlines():
        stripped = line.strip()
        if not stripped.startswith("|") or "---" in stripped:
            continue
        cells = [cell.strip() for cell in stripped.strip("|").split("|")]
        if not cells:
            continue
        code = cells[0]
        if code.lower() == "code":
            continue
        if re.fullmatch(r"[A-Z0-9]+", code):
            codes.add(code)
    return codes


def csv_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def extract_bls_import(
    *,
    workbook_path: Path = DEFAULT_WORKBOOK_PATH,
    nutrient_defs_path: Path = DEFAULT_NUTRIENT_DEFS_PATH,
    output_dir: Path = DEFAULT_OUTPUT_DIR,
    limit: int | None = None,
) -> ImportResult:
    if limit is not None and limit <= 0:
        raise ValueError("limit must be positive when provided")
    if not workbook_path.exists():
        raise FileNotFoundError(f"BLS workbook is missing: {workbook_path}")
    if not nutrient_defs_path.exists():
        raise FileNotFoundError(f"nutrient_defs candidate is missing: {nutrient_defs_path}")

    supported_codes = parse_supported_nutrient_codes(nutrient_defs_path.read_text(encoding="utf-8"))
    if not supported_codes:
        raise ValueError(f"no supported nutrient codes parsed from {nutrient_defs_path}")

    output_dir.mkdir(parents=True, exist_ok=True)
    foods_path = output_dir / "foods.csv"
    food_nutrients_path = output_dir / "food_nutrients.csv"

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        header = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        nutrient_columns: list[tuple[int, str]] = []
        unsupported_codes: set[str] = set()
        for column_index in range(4, len(header) + 1, 3):
            code = extract_nutrient_code(header[column_index - 1])
            if code is None:
                continue
            if code in supported_codes:
                nutrient_columns.append((column_index, code))
            else:
                unsupported_codes.add(code)

        food_count = 0
        food_nutrient_count = 0
        missing_value_count = 0
        used_nutrient_codes: set[str] = set()
        sample_foods: list[tuple[str, str, str]] = []

        with foods_path.open("w", encoding="utf-8", newline="") as foods_file, food_nutrients_path.open("w", encoding="utf-8", newline="") as nutrients_file:
            foods_writer = csv.writer(foods_file, lineterminator="\n")
            nutrients_writer = csv.writer(nutrients_file, lineterminator="\n")
            foods_writer.writerow(["bls_code", "name_de", "name_en", "name_th", "name_display"])
            nutrients_writer.writerow(["bls_code", "nutrient_code", "value", "data_source"])

            for row in sheet.iter_rows(min_row=2, values_only=True):
                bls_code = csv_text(row[0])
                name_de = csv_text(row[1])
                name_en = csv_text(row[2])
                if not bls_code or not name_de:
                    continue

                foods_writer.writerow([bls_code, name_de, name_en, "", name_de])
                food_count += 1
                if len(sample_foods) < 10:
                    sample_foods.append((bls_code, name_de, name_en))

                for column_index, nutrient_code in nutrient_columns:
                    value = row[column_index - 1] if column_index - 1 < len(row) else None
                    parsed = parse_bls_decimal(value)
                    if parsed is None:
                        missing_value_count += 1
                        continue
                    nutrients_writer.writerow([bls_code, nutrient_code, format(parsed, "f"), SOURCE_LABEL])
                    food_nutrient_count += 1
                    used_nutrient_codes.add(nutrient_code)

                if limit is not None and food_count >= limit:
                    break

        sheet_name = sheet.title
    finally:
        workbook.close()

    if food_count == 0:
        raise ValueError("no source-backed foods were extracted")
    if food_nutrient_count == 0:
        raise ValueError("no source-backed food nutrient values were extracted")

    return ImportResult(
        source_workbook=workbook_path,
        nutrient_defs_path=nutrient_defs_path,
        output_dir=output_dir,
        sheet_name=sheet_name,
        food_count=food_count,
        food_nutrient_count=food_nutrient_count,
        supported_nutrient_count=len(supported_codes),
        distinct_nutrient_count=len(used_nutrient_codes),
        missing_value_count=missing_value_count,
        unsupported_nutrient_codes=unsupported_codes,
        sample_foods=sample_foods,
    )


def build_apply_sql(container_import_dir: str = CONTAINER_IMPORT_DIR) -> str:
    return f"""-- STATUS: LOCAL_ONLY_BLS_IMPORT_EXPANSION
-- PURPOSE: Apply deterministic UTF-8 CSV artifacts generated from the approved local BLS workbook.
-- AUTHORIZATION BOUNDARY:
--   - Local Supabase/Test DB only
--   - No DEV/LIVE
--   - No Supabase Cloud
--   - No source workbook commit
--   - No unsupported food or nutrient values
--   - No RDA value changes

begin;

create temporary table stage_foods (
  bls_code text not null,
  name_de text not null,
  name_en text,
  name_th text,
  name_display text
) on commit drop;

create temporary table stage_food_nutrients (
  bls_code text not null,
  nutrient_code text not null,
  value numeric(12,5) not null,
  data_source text not null
) on commit drop;

\\copy stage_foods (bls_code, name_de, name_en, name_th, name_display) from '{container_import_dir}/foods.csv' with (format csv, header true, encoding 'UTF8')

\\copy stage_food_nutrients (bls_code, nutrient_code, value, data_source) from '{container_import_dir}/food_nutrients.csv' with (format csv, header true, encoding 'UTF8')

do $$
declare
  missing_nutrient_count integer;
begin
  select count(*) into missing_nutrient_count
  from stage_food_nutrients staged
  left join nutrition.nutrient_defs defs
    on defs.code = staged.nutrient_code
  where defs.code is null;

  if missing_nutrient_count <> 0 then
    raise exception 'staged food_nutrients contain % missing nutrient_defs targets', missing_nutrient_count;
  end if;
end
$$;

insert into nutrition.foods (bls_code, name_de, name_en, name_th, name_display)
select bls_code, name_de, name_en, coalesce(name_th, ''), name_display
from stage_foods
on conflict (bls_code) do update set
  name_de = excluded.name_de,
  name_en = excluded.name_en,
  name_th = excluded.name_th,
  name_display = excluded.name_display,
  updated_at = now();

insert into nutrition.food_nutrients (food_id, nutrient_code, value, data_source)
select foods.id, staged.nutrient_code, staged.value, staged.data_source
from stage_food_nutrients staged
join nutrition.foods foods
  on foods.bls_code = staged.bls_code
join nutrition.nutrient_defs defs
  on defs.code = staged.nutrient_code
on conflict (food_id, nutrient_code) do update set
  value = excluded.value,
  data_source = excluded.data_source;

analyze nutrition.foods;
analyze nutrition.food_nutrients;

commit;
"""


def build_report(result: ImportResult) -> str:
    unsupported = ", ".join(sorted(result.unsupported_nutrient_codes)[:10]) or "none"
    scope = "full" if result.food_count >= 7140 else "subset"
    sample_rows = "\n".join(
        f"| {code} | {name_de} | {name_en} |"
        for code, name_de, name_en in result.sample_foods
    )
    return f"""# P1-005 Local BLS Import Expansion Report

Status: LOCAL_ONLY_IMPORT_EXPANSION_CANDIDATE

## Purpose

Expand the local Nutrition food foundation with deterministic source-backed BLS
foods and nutrient values so the next local product step can build real Food
Search.

## Source

- Source workbook: `{result.source_workbook.as_posix()}`
- Sheet: `{result.sheet_name}`
- Nutrient code allowlist: `{result.nutrient_defs_path.as_posix()}`
- Extraction scope: `{scope}`
- Food rows extracted: {result.food_count}
- Food nutrient rows extracted: {result.food_nutrient_count}
- Supported nutrient definitions: {result.supported_nutrient_count}
- Distinct nutrient codes used: {result.distinct_nutrient_count}
- Missing/trace/non-numeric nutrient cells skipped: {result.missing_value_count}
- Unsupported nutrient header mappings: {len(result.unsupported_nutrient_codes)}
- Unsupported examples: {unsupported}

## Local artifacts

The UTF-8 CSV files are generated under `{result.output_dir.as_posix()}` and are
local runtime artifacts, not committed BLS data:

- `foods.csv`
- `food_nutrients.csv`
- `apply-local.sql`

## Sample foods

| bls_code | name_de | name_en |
|---|---|---|
{sample_rows}

## Boundaries

- Local Supabase/Test DB only.
- No DEV/LIVE action.
- No Supabase Cloud command.
- No source workbook commit.
- No unsupported food values.
- No unsupported nutrient values.
- No RDA value changes.
- No schema change.

## Validation queries

```sql
select count(*) from nutrition.foods;
select count(*) from nutrition.food_nutrients;
select count(*)
from nutrition.food_nutrients fn
left join nutrition.nutrient_defs nd on nd.code = fn.nutrient_code
where nd.code is null;
select count(*)
from nutrition.food_nutrients fn
left join nutrition.foods foods on foods.id = fn.food_id
where foods.id is null;
```
"""


def write_import_artifacts(result: ImportResult, report_path: Path = DEFAULT_REPORT_PATH) -> None:
    (result.output_dir / "apply-local.sql").write_text(build_apply_sql(), encoding="utf-8", newline="\n")
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(build_report(result), encoding="utf-8", newline="\n")


def main() -> int:
    parser = argparse.ArgumentParser(description="Generate deterministic local-only BLS import expansion artifacts.")
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK_PATH)
    parser.add_argument("--nutrient-defs", type=Path, default=DEFAULT_NUTRIENT_DEFS_PATH)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--report", type=Path, default=DEFAULT_REPORT_PATH)
    parser.add_argument("--limit", type=int, default=None)
    args = parser.parse_args()

    result = extract_bls_import(
        workbook_path=args.workbook,
        nutrient_defs_path=args.nutrient_defs,
        output_dir=args.output_dir,
        limit=args.limit,
    )
    write_import_artifacts(result, args.report)
    print(f"scope={'full' if args.limit is None else 'subset'}")
    print(f"foods={result.food_count}")
    print(f"food_nutrients={result.food_nutrient_count}")
    print(f"missing_values={result.missing_value_count}")
    print(f"unsupported_mappings={len(result.unsupported_nutrient_codes)}")
    print(f"output_dir={args.output_dir}")
    print(f"report={args.report}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
